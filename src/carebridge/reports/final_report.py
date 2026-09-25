"""Assemble the capstone report from computed tables; fail if results are missing."""
from __future__ import annotations

import hashlib
import json
from datetime import date
from importlib.metadata import version

import numpy as np
import pandas as pd
from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt

from carebridge.config import RAW, REPORTS, ROOT

SECTIONS = ["Abstract", "Executive summary", "1. Introduction", "2. Research questions",
            "3. Data preparation", "4. Methodology", "5. Analysis",
            "6. Recommendations", "7. Challenges and limitations", "8. References", "Appendix"]


def read(name):
    return pd.read_csv(REPORTS / "tables" / f"{name}.csv")


def display(value, column=""):
    if pd.isna(value):
        return "—"
    if column in {"p_value", "p_adj", "p_adj_bh"} and value == 0:
        return "<1e-300"
    if isinstance(value, (float, np.floating)):
        if value != 0 and abs(value) < 0.0001:
            return f"{value:.2e}"
        return f"{value:.4f}"
    return str(value)


class Report:
    def __init__(self):
        self.document = Document()
        self.document.styles["Normal"].font.name = "Calibri"
        self.document.styles["Normal"].font.size = Pt(10)
        for section in self.document.sections:
            section.left_margin = section.right_margin = Inches(0.75)
            footer = section.footer.paragraphs[0]
            footer.alignment = 2
            footer.add_run("CareBridge | ")
            page = OxmlElement("w:fldSimple")
            page.set(qn("w:instr"), "PAGE")
            footer._p.append(page)
        self.lines = []
        self.document.add_heading("CareBridge: Diabetes Readmission and Cardiometabolic Comorbidity", 0)
        self.document.add_paragraph(f"Antoine Ward | B.S. Data Science capstone | {date.today()}")
        self.lines.extend(["# CareBridge: Diabetes Readmission and Cardiometabolic Comorbidity",
                           f"Antoine Ward | B.S. Data Science capstone | {date.today()}", ""])
        self.heading("Contents")
        for title in SECTIONS:
            self.text(title)

    def heading(self, title, level=1):
        self.document.add_heading(title, level)
        self.lines.extend(["#" * (level + 1) + " " + title, ""])

    def text(self, text):
        self.document.add_paragraph(text)
        self.lines.extend([text, ""])

    def table(self, frame, caption):
        self.text(caption)
        table = self.document.add_table(rows=1, cols=len(frame.columns))
        table.style = "Light Shading Accent 1"
        repeat_header = OxmlElement("w:tblHeader")
        table.rows[0]._tr.get_or_add_trPr().append(repeat_header)
        for cell, name in zip(table.rows[0].cells, frame.columns):
            cell.text = str(name)
        for values in frame.itertuples(index=False, name=None):
            for cell, value, column in zip(table.add_row().cells, values, frame.columns):
                cell.text = display(value, column)
        self.lines.append("| " + " | ".join(map(str, frame.columns)) + " |")
        self.lines.append("| " + " | ".join(["---"] * len(frame.columns)) + " |")
        for values in frame.itertuples(index=False, name=None):
            self.lines.append("| " + " | ".join(display(v, c).replace("|", "/")
                                                 for v, c in zip(values, frame.columns)) + " |")
        self.lines.append("")

    def figure(self, name, caption):
        path = REPORTS / "figures" / name
        if not path.exists():
            raise FileNotFoundError(path)
        self.document.add_picture(str(path), width=Inches(6))
        self.document.add_paragraph(caption, style="Caption")
        self.lines.extend([f"![{caption}](figures/{name})", ""])

    def save(self):
        self.document.save(REPORTS / "CareBridge_Final_Report.docx")
        (REPORTS / "CareBridge_Final_Report.md").write_text("\n".join(self.lines))


def manifest():
    def digest(path):
        result = hashlib.sha256()
        with path.open("rb") as stream:
            for chunk in iter(lambda: stream.read(1024 * 1024), b""):
                result.update(chunk)
        return result.hexdigest()
    paths = sorted((REPORTS / "tables").glob("*.csv"))
    sources = sorted(RAW.rglob("*.csv"))
    code = sorted((ROOT / "src").rglob("*.py")) + sorted((ROOT / "sql").rglob("*.sql"))
    metadata = {
        "generated_date": str(date.today()),
        "versions": {name: version(name) for name in
                     ["numpy", "pandas", "scikit-learn", "statsmodels", "xgboost", "shap",
                      "tensorflow", "ml-dtypes", "duckdb"]},
        "raw_sha256": {str(path.relative_to(ROOT)): digest(path) for path in sources},
        "table_sha256": {str(path.relative_to(ROOT)): digest(path) for path in paths},
        "code_sha256": {str(path.relative_to(ROOT)): digest(path) for path in code},
    }
    (REPORTS / "analysis_manifest.json").write_text(json.dumps(metadata, indent=2) + "\n")


def main():
    discrimination = read("rq1_model_discrimination").set_index("model")
    comparison = read("rq1_model_comparisons").set_index("comparison")
    calibration = read("rq1_calibration").drop_duplicates("model").set_index("model")
    leakage = read("rq1_leakage_sensitivity")
    leak_difference = read("rq1_leakage_comparison").iloc[0]
    shap = read("rq1_shap_global_importance")
    gaps = read("rq1_fairness_gaps")
    rq2 = read("rq2_discrimination")
    nested = read("rq2_nested_lrt")
    comorbidity = read("rq2_comorbidity")
    mlp = read("rq1_mlp_comparison")
    mlp_difference = read("rq1_mlp_difference").iloc[0]
    profiles = read("rq1_phenotype_profiles")
    h1b = read("h1b_adjusted_odds")
    h1e = read("h1e_poisson_vs_negbin").iloc[0]
    n = int(calibration.iloc[0]["n"])
    if set(leakage["n"]) != {n} or int(h1e["n"]) != n:
        raise ValueError("RQ1 artifacts were built from different cohort sizes; rerun analysis")
    gb = discrimination.loc["gradient_boosting"]
    prev = discrimination.loc["prevalence", "pr_auc"]
    effect = comparison.loc["gradient_boosting - logistic"]
    report = Report()
    report.heading("Abstract")
    report.text(
        f"This study examined readmission prediction in {n:,} patients with diabetes and "
        f"cardiometabolic associations in {int(rq2.iloc[0]['n']):,} BRFSS respondents. "
        f"Patient-grouped cross-validation yielded XGBoost average precision (AP) {gb.pr_auc:.4f} "
        f"(95% interval {gb.ci_low:.4f}–{gb.ci_high:.4f}) at {prev:.2%} event prevalence. "
        f"Adding a deliberately unavailable future encounter count increased AP by "
        f"{leak_difference.pr_auc_difference:.4f}, illustrating the impact of information leakage. "
        "SHAP described the fitted model; threshold audits quantified demographic differences. "
        "Three parallel survey models, nested likelihood tests, and adjusted comorbidity models "
        "characterized cross-sectional associations. An MLP and outcome-blind clustering extended "
        "the analysis. Results support further validation, with substantial limits from historical "
        "data, sample selection, missing survey-design fields, and observational measurement.")
    report.heading("Executive summary")
    report.text(
        f"XGBoost exceeded the utilization-history rule (AP "
        f"{discrimination.loc['rule_baseline', 'pr_auc']:.4f}) and logistic regression "
        f"({discrimination.loc['logistic', 'pr_auc']:.4f}). Its lift over prevalence was "
        f"{gb.lift:.2f}. The raw probabilities were audited for calibration but were not "
        "recalibrated. Group error rates vary at a shared cutoff, so aggregate accuracy cannot "
        "establish equitable performance. The future-count model is retained solely as a diagnostic. "
        "The BRFSS outcomes measure existing diagnosed conditions, and cardiac disease is a combined "
        "heart disease/angina/heart attack indicator. All survey percentages are unweighted sample "
        "statistics. The deliverable includes code, tables, figures, this report, and a reproducible runner.")
    report.heading("1. Introduction")
    report.text(
        "CareBridge connects two analytic questions: identifying patterns associated with early "
        "readmission after a diabetes hospitalization, and comparing correlates of diabetes, stroke, "
        "and cardiac disease in a community survey. The sources differ in time, population, and unit "
        "of observation, so they are analyzed separately. The hospital dataset covers 1999–2008 [1]; "
        "the survey extract derives from 2015 BRFSS [2, 3]. Neither analysis estimates intervention effects.")
    report.heading("2. Research questions")
    report.text(
        "RQ1: Which patient, encounter, utilization, and medication-management factors predict "
        "30-day readmission, and how do discrimination, calibration, and subgroup error rates compare? "
        "The proposal's intent concerns unplanned readmission, but the supplied target does not "
        "distinguish planned readmissions. The available predictors support scoring at discharge.")
    report.text(
        "RQ2: Which behavioral, clinical, and socioeconomic indicators predict each survey condition, "
        "how do their rankings differ, and which existing diagnoses are most strongly associated "
        "with another condition? This is concurrent disease classification, not a prospective incidence model.")
    report.text(
        "The original verbatim H2a–H2d statements were unavailable in the repository. This report uses "
        "the working definitions in analysis_specification.md: H2a tests the clinical/behavioral block; "
        "H2b tests the socioeconomic/access block; H2c describes ranking differences; H2d examines "
        "adjusted comorbidity. These are exploratory operational definitions, not a claim of preregistration. "
        "H1d has no supplied parity margin, so the audit cannot declare equivalence or certify fairness.")
    report.heading("3. Data preparation")
    report.text(
        "Bronze preserves raw strings and lineage. Silver casts types, resolves sentinels, and applies "
        "eligibility rules. Gold groups primary diagnoses, derives encounter order, and selects each "
        "patient's first retained encounter. Encounter ID is an ordering proxy because admission dates "
        "are unavailable. Expired discharges are outside the population at risk for subsequent readmission; "
        "hospice is retained. Patients with apparent post-expired encounters are excluded as unresolved "
        "records. These are eligibility and data-quality decisions; the deliberate future-count feature "
        "is a separate temporal-leakage experiment. Stable patient/encounter sorting makes CV input order repeatable.")
    report.text(
        "Medication adjustment is defined as any individual drug marked Up or Down. The original change "
        "flag also occurs with No/Steady drug entries and therefore does not isolate this dose-adjustment "
        "construct. This finding does not establish that no medication initiation or other treatment change "
        "occurred. Features include medication counts, service counts per day, prior utilization, and "
        "an Isolation Forest anomaly score. Race is reserved for audit; age and recorded gender remain predictors.")
    report.text(
        "BRFSS uses the unrebalanced 253,680-row binary extract, not the 50/50 variant. All three outcome "
        "columns are excluded from the shared predictive feature set. Other diagnosis indicators enter "
        "only the separate, explicitly cross-sectional comorbidity models. Survey weights, strata, and "
        "primary sampling units are absent from this mirror, which limits population inference [2, 3].")
    report.heading("4. Methodology")
    report.heading("4.1 Readmission prediction and uncertainty", 2)
    report.text(
        "The utilization rule, logistic regression, linear SVM, and XGBoost were compared on five "
        "StratifiedGroupKFold partitions (seed 42). Anomaly detection, scaling, and categorical encoding "
        "were fitted inside each training fold. One patient contributes one primary row. AP is the "
        "reported PR-AUC convention, with ROC-AUC secondary. Paired percentile intervals use 1,000 "
        "patient bootstrap resamples of fixed OOF scores. They do not include model-refit or selection "
        "uncertainty. No cutoff was optimized on these outcomes.")
    report.heading("4.2 Calibration, explanation, and sensitivity", 2)
    report.text(
        "Raw OOF probability calibration uses Brier score, a joint logistic intercept/slope fit, and "
        "ten quantile bins. The scikit-learn return order is observed fraction followed by mean prediction "
        "[4]. SHAP explains the full-cohort XGBoost refit using 5,000 sampled observations, an explicit "
        "100-row background, and interventional probability output [5]. Base value plus contributions "
        "must reconstruct each prediction within numerical tolerance. Source-level importance sums "
        "signed one-hot contributions per patient before taking absolute magnitudes. This is "
        "descriptive model attribution, not causal inference or a new performance estimate. Correlated "
        "and derived features complicate individual attribution.")
    report.text(
        "The sensitivity experiment adds patient_encounter_count to an otherwise identical pipeline "
        "on identical materialized folds. This count includes later retained encounters and is unavailable "
        "at the index discharge. The primary feature allowlist continues to exclude it. Subgroup audits "
        "report TPR and FPR at 0.05, 0.10, 0.15, 0.20, and 0.25 with pointwise Wilson 95% intervals. "
        "Zero denominators produce undefined rates. Descriptive max–min gaps exclude denominators below 20; "
        "all group counts remain available. The displayed 0.10 cutoff is illustrative.")
    report.heading("4.3 Survey models and hypothesis tests", 2)
    report.text(
        "Three common logistic specifications use five shared folds stratified by the joint outcome "
        "pattern. Age, general health, education, and income bands are one-hot encoded; continuous "
        "predictors are standardized inside folds. AP intervals use 300 respondent-row resamples. "
        "Held-out permutation importance uses 5,000 validation rows per fold and three repeats; the "
        "reported standard deviation is repeat/fold variability, not a confidence interval. Separate "
        "unpenalized logistic association models estimate odds ratios and nested likelihood-ratio tests "
        "on identical rows. Benjamini–Hochberg correction is applied separately to nested tests, "
        "comorbidity terms, and the coefficient table. Standard errors assume independent records and "
        "do not account for the unavailable survey design.")
    report.heading("4.4 Neural benchmark and phenotyping", 2)
    report.text(
        "The Keras benchmark uses 64 and 32 ReLU units, dropout 0.2, a sigmoid output, Adam at 0.001, "
        "batch size 512, and exactly 20 training epochs per fold. Test patients never influence "
        "preprocessing or stopping. It is compared with XGBoost using paired OOF AP. Exploratory "
        "K-means phenotypes use 12 utilization inputs after log1p and standardization; outcomes are "
        "excluded. The cluster count (2–6) is selected by silhouette on a fixed sample, then refitted "
        "on the full cohort. A second seed provides an adjusted Rand agreement check. Post-fit event "
        "rates describe the clusters and do not independently validate them.")
    report.heading("5. Analysis")
    report.heading("5.1 Model discrimination and existing hypotheses", 2)
    report.table(discrimination.reset_index()[["model", "pr_auc", "ci_low", "ci_high", "roc_auc"]],
                 "Table 1. RQ1 OOF discrimination; PR-AUC denotes average precision.")
    report.text(
        f"H1c: the paired XGBoost-minus-logistic bootstrap mean AP difference was {effect.mean_diff:.4f} "
        f"({effect.ci_low:.4f} to {effect.ci_high:.4f}). The difference is positive in the reported interval. "
        "This comparison remains conditional on the chosen feature set, models, and cohort.")
    interaction = h1b[h1b["term"].eq("Both (interaction)")].iloc[0]
    report.text(
        f"H1b: the dose-adjustment interaction OR was {interaction.odds_ratio:.3f} "
        f"({interaction.ci_low:.3f}–{interaction.ci_high:.3f}); p={interaction.p_value:.4g}. "
        "An interaction coefficient modifies the product of main-effect odds ratios; it is not "
        "the combined group's standalone odds ratio. These observational estimates do not identify "
        "a treatment benefit. "
        f"H1e: length-of-stay variance/mean was {h1e.dispersion_ratio:.4f}; the boundary-corrected "
        f"Poisson-versus-negative-binomial LR statistic was {h1e.lr_statistic:.1f}. The fitted negative "
        "binomial improves likelihood, but the observed stay range is restricted to 1–14 days and "
        "untruncated distribution fits remain approximations.")
    report.heading("5.2 Calibration and SHAP", 2)
    report.table(calibration.reset_index()[["model", "brier_score", "calibration_intercept",
                                            "calibration_slope"]], "Table 2. Raw OOF calibration.")
    report.figure("rq1_calibration.png", "Figure 1. Correctly oriented predicted versus observed risk.")
    report.table(shap.head(10)[["source_feature", "mean_abs_shap", "importance_pct"]],
                 "Table 3. Ten highest source-level SHAP magnitudes; units are probability contributions.")
    report.figure("rq1_shap_global_importance.png", "Figure 2. Source-level SHAP importance.")
    report.figure("rq1_shap_beeswarm.png", "Figure 3. Encoded-feature SHAP contributions in the sampled cohort.")
    report.heading("5.3 Leakage sensitivity", 2)
    report.table(leakage[["scenario", "pr_auc", "ci_low", "ci_high", "roc_auc"]],
                 "Table 4. Identical-fold comparison with and without future encounter count.")
    report.text(
        f"The observed AP increase was {leak_difference.pr_auc_difference:.4f}, with paired 95% "
        f"interval {leak_difference.ci_low:.4f}–{leak_difference.ci_high:.4f}. This is apparent "
        "performance obtained using unavailable information; it must not be reported as a valid model improvement.")
    report.figure("rq1_leakage_sensitivity.png", "Figure 4. Performance inflation from future information.")
    report.heading("5.4 H1d subgroup audit", 2)
    report.table(gaps[gaps["threshold"].eq(0.10)][["attribute", "metric", "eligible_groups", "max_minus_min"]],
                 "Table 5. Descriptive subgroup ranges at the illustrative 0.10 threshold.")
    report.figure("rq1_fairness.png", "Figure 5. Subgroup TPR/FPR and pointwise Wilson intervals.")
    report.text(
        "Differences in observed error rates require investigation and cannot be dismissed because "
        "race was excluded from predictors. Small groups have wider uncertainty. Marginal intervals "
        "and max–min ranges do not constitute a simultaneous test or an equivalence assessment. "
        "H1d is addressed descriptively; a binary fairness conclusion is not identified without a "
        "prespecified tolerance, deployment context, and independent validation.")
    report.heading("5.5 RQ2 and working H2a–H2d", 2)
    report.table(rq2[["outcome", "prevalence", "pr_auc", "ci_low", "ci_high", "roc_auc"]],
                 "Table 6. Three survey outcome models; sample-level OOF performance.")
    report.table(nested[["outcome", "comparison", "lr_statistic", "df", "p_adj_bh"]],
                 "Table 7. Working H2a/H2b nested likelihood tests.")
    report.text(
        f"At BH-adjusted alpha 0.05, {(nested.p_adj_bh < 0.05).sum()} of {len(nested)} block-addition "
        "tests reject their reduced specifications. Large samples can detect small incremental effects; "
        "these p-values do not measure predictive or practical value. H2c is described by held-out "
        "permutation rankings and their Spearman agreement below, without a formal equality claim.")
    report.figure("rq2_predictor_importance.png", "Figure 6. Outcome-specific held-out permutation AP losses.")
    for outcome, sub in read("rq2_importance").groupby("outcome"):
        leaders = ", ".join(sub.nlargest(3, "mean_ap_decrease")["feature"])
        report.text(f"The three largest held-out permutation AP losses for {outcome} were "
                    f"{leaders}. These ranks reflect the fitted specification and predictor dependence.")
    report.table(read("rq2_rank_agreement"), "Table 8. Descriptive agreement between outcome predictor rankings.")
    report.table(comorbidity[["outcome", "existing_condition", "conditional_prevalence",
                              "adjusted_odds_ratio", "or_ci_low", "or_ci_high"]],
                 "Table 9. Working H2d: directed comorbidity associations adjusted for shared covariates and the third diagnosis.")
    for outcome, sub in comorbidity.groupby("outcome"):
        best = sub.loc[sub["adjusted_odds_ratio"].idxmax()]
        report.text(f"For {outcome}, {best.existing_condition} had the largest adjusted odds-ratio "
                    f"point estimate ({best.adjusted_odds_ratio:.3f}; {best.or_ci_low:.3f}–{best.or_ci_high:.3f}). "
                    "Ranking point estimates does not establish that competing associations differ significantly.")
    report.heading("5.6 MLP and exploratory phenotypes", 2)
    report.table(mlp, "Table 10. Exploratory neural-network benchmark.")
    report.text(f"MLP minus XGBoost AP was {mlp_difference.pr_auc_difference:.4f}, with paired interval "
                f"{mlp_difference.ci_low:.4f}–{mlp_difference.ci_high:.4f}. No neural hyperparameter "
                "search was performed. Model selection on this cohort requires a separate external evaluation.")
    report.table(profiles[["cluster", "n", "event_rate", "ci_low", "ci_high", "seed_stability_ari"]],
                 "Table 11. Phenotype size, post-fit event rate, and agreement across two seeds.")
    report.figure("rq1_phenotypes.png", "Figure 7. Utilization phenotype profiles; cluster numbers are arbitrary.")
    report.heading("6. Recommendations")
    report.text(
        "Retain the leakage-safe feature contract and validate predictor availability at discharge. "
        "Evaluate the frozen pipeline on newer data from independent institutions before considering "
        "operational use. Fit any probability recalibration within training folds or on a separate "
        "calibration set. Select operating points using a stated workload and error-cost objective, "
        "then assess uncertainty on independent patients. Reassess subgroup errors with adequate "
        "event counts and prespecified disparity tolerances. Obtain the original weighted BRFSS "
        "design variables for population inference. Treat comorbidity and phenotype results as "
        "hypothesis-generating associations, not individual care recommendations.")
    report.heading("7. Challenges and limitations")
    report.text(
        "The principal challenges were temporal leakage, ambiguous medication coding, sparse prior "
        "utilization counts, and unequal subgroup precision. Calibration plotting originally swapped "
        "observed and predicted outputs; a regression test now verifies the exported columns. SHAP "
        "aggregation now preserves signed dummy contributions before calculating magnitudes. Stable "
        "cohort ordering and pinned dependencies reduce reproducibility drift. Patient-level predictions "
        "and explanations stay in the ignored local data directory; published outputs are aggregate.")
    report.text(
        "Residual limits include historical and selected hospital data, absent admission dates and "
        "planned-readmission flags, possible missed care outside participating hospitals, use of "
        "encounter ID as chronology, observational confounding, and selection of the first retained "
        "encounter. Goodness-of-fit rejection is expected with large samples; continuous fits to "
        "integer length of stay are descriptive comparisons rather than proof of an underlying continuous "
        "distribution. BRFSS is self-reported, cross-sectional, and unweighted here. Symptoms, general "
        "health, and diagnoses can be jointly determined. Survey-response patterns cannot identify "
        "unique people or households in this extract. Bootstrap uncertainty is conditional on existing "
        "OOF predictions. SHAP background choice, correlated predictors, permutation dependence, and "
        "cluster selection can change rankings. The original formal H2 statements and institutional "
        "submission template must be reconciled with the separate course materials before submission.")
    report.heading("8. References")
    references = [
        ("[1] Clore, Cios, DeShazo, and Strack (2014). Diabetes 130-US Hospitals for Years 1999–2008. "
         "UCI Machine Learning Repository. https://doi.org/10.24432/C5230J"),
        ("[2] CDC. 2015 BRFSS Survey Data and Documentation. "
         "https://www.cdc.gov/brfss/annual_data/annual_2015.html"),
        ("[3] CDC. BRFSS 2015 Codebook. "
         "https://www.cdc.gov/brfss/annual_data/2015/pdf/codebook15_llcp.pdf"),
        ("[4] Scikit-learn 1.5. calibration_curve API documentation. "
         "https://scikit-learn.org/1.5/modules/generated/sklearn.calibration.calibration_curve.html"),
        ("[5] SHAP. TreeExplainer API documentation. "
         "https://shap.readthedocs.io/en/stable/generated/shap.TreeExplainer.html"),
        ("[6] Strack et al. (2014). Impact of HbA1c Measurement on Hospital Readmission Rates: "
         "Analysis of 70,000 Clinical Database Patient Records. BioMed Research International, 781670. "
         "https://doi.org/10.1155/2014/781670"),
        ("[7] Dataset mirrors: https://www.kaggle.com/datasets/brandao/diabetes and "
         "https://www.kaggle.com/datasets/alexteboul/diabetes-health-indicators-dataset. "
         "Original source documentation takes precedence over mirror descriptions."),
    ]
    for reference in references:
        report.text(reference)
    report.heading("Appendix")
    report.heading("A. Reproduction and artifact provenance", 2)
    report.text(
        "Install Python 3.12 and requirements.txt, install the local package, configure Kaggle "
        "credentials, run python -m carebridge.ingest.download, and run python -m carebridge.run_analysis. "
        "The runner executes sequentially with bounded thread counts. Each analysis also runs as its "
        "own module. analysis_manifest.json records package versions and SHA-256 hashes of input "
        "files and result tables. analysis_specification.md records feature, uncertainty, and "
        "hypothesis-label conventions. The table of contents is assembled from this report's fixed "
        "section list; it contains section titles rather than renderer-dependent page numbers.")
    appendix_tables = [
        ("rq1_model_comparisons", None), ("rq1_threshold_analysis", ["rule", "threshold", "flagged_n", "sensitivity_recall", "ppv_precision", "f1"]),
        ("rq1_calibration", ["model", "bin", "mean_predicted_probability", "observed_event_rate"]),
        ("h1b_adjusted_odds", ["term", "odds_ratio", "ci_low", "ci_high", "p_value"]),
        ("rq1_hypothesis_tests", ["test", "effect_size", "p_adj", "n"]),
        ("rq2_importance", ["outcome", "feature", "mean_ap_decrease", "repeat_sd", "rank"]),
        ("rq2_adjusted_odds", ["outcome", "term", "odds_ratio", "ci_low", "ci_high", "p_adj_bh"]),
        ("rq1_phenotype_selection", None),
        ("rq1_shap_global_importance", None),
        ("rq1_shap_encoded_importance", None),
        ("rq2_comorbidity", ["outcome", "existing_condition", "exposed_n", "cooccurring_n",
                             "conditional_prevalence", "prevalence_ci_low", "prevalence_ci_high"]),
        ("rq2_comorbidity", ["outcome", "existing_condition", "prevalence_lift", "p_adj_bh"]),
        ("rq2_discrimination", ["outcome", "n", "events", "brier_score"]),
    ]
    for i, (name, columns) in enumerate(appendix_tables, start=1):
        table = read(name)
        report.table(table if columns is None else table[columns], f"Appendix table A{i}: {name}.csv")
    for metric, denominator in [("tpr", "events"), ("fpr", "non_events")]:
        table = read("rq1_fairness")[["attribute", "group", "threshold", denominator,
                                      metric, metric + "_ci_low", metric + "_ci_high"]]
        report.table(table, f"Subgroup {metric.upper()} at every prespecified threshold (rq1_fairness.csv).")
    profile_means = profiles.melt(id_vars="cluster", value_vars=[c for c in profiles if c.startswith("mean_")],
                                  var_name="feature", value_name="mean")
    report.table(profile_means, "Raw-scale utilization means by phenotype (rq1_phenotype_profiles.csv).")
    report.heading("B. Artifact inventory", 2)
    for path in sorted((REPORTS / "tables").glob("*.csv")):
        report.text(str(path.relative_to(REPORTS)))
    manifest()
    report.save()
    print("Wrote reports/CareBridge_Final_Report.md, .docx, and analysis_manifest.json")


if __name__ == "__main__":
    main()
